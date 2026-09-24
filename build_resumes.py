"""Build role-specific resume variants (DOCX + PDF) from resumes/content.yaml.

  python build_resumes.py                 # all variants -> resumes/<variant>.docx / .pdf
  python build_resumes.py ai_ml_engineer  # just one
  python build_resumes.py --no-pdf        # DOCX only

Layout: single column, standard section headings, real text (no tables / images / columns),
so applicant-tracking systems can parse it. US Letter, Times New Roman, one page.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
RES_DIR = ROOT / "resumes"
FONT = "Times New Roman"
BODY_PT = 10.5
TEXT_WIDTH_IN = 7.0  # 8.5in page - 2 x 0.75in margins
LINK_COLOR = RGBColor(0x1F, 0x4E, 0x79)


# ------------------------------------------------------------------ low-level helpers
def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def spacing(paragraph, before=0, after=0, line=1.0):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_hyperlink(paragraph, url: str, text: str, size: float):
    """Clickable link (python-docx has no public API for this)."""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)
    rpr.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def bottom_border(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    ppr.append(borders)


# ------------------------------------------------------------------ building blocks
class ResumeWriter:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.top_margin = sec.bottom_margin = Inches(0.6)
        sec.left_margin = sec.right_margin = Inches(0.75)
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(BODY_PT)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        bullet = self.doc.styles["List Bullet"]
        bullet.font.name = FONT
        bullet.font.size = Pt(BODY_PT)

    def para(self, align=None, before=0, after=0):
        p = self.doc.add_paragraph()
        if align:
            p.alignment = align
        spacing(p, before, after)
        return p

    def text(self, p, text, **fmt):
        run = p.add_run(text)
        set_run_font(run, **fmt)
        return run

    def header(self, h: dict, headline: str):
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER)
        self.text(p, h["name"], size=17, bold=True)
        if headline:
            p = self.para(WD_ALIGN_PARAGRAPH.CENTER, before=1)
            self.text(p, headline, size=10.5, italic=True)
        p = self.para(WD_ALIGN_PARAGRAPH.CENTER, before=2)
        parts = [h.get("phone"), h.get("email")]
        self.text(p, "  |  ".join(x for x in parts if x), size=10)
        for link in h.get("links") or []:
            self.text(p, "  |  ", size=10)
            add_hyperlink(p, link["url"], link["label"], 10)
        if h.get("location"):
            p = self.para(WD_ALIGN_PARAGRAPH.CENTER, before=1)
            self.text(p, h["location"], size=10)

    def section(self, title: str):
        p = self.para(before=7, after=3)
        self.text(p, title.upper(), size=11, bold=True)
        bottom_border(p)

    def entry(self, left: str, right: str = "", sub: str = "", before: float = 3):
        """Bold title with right-aligned date, optional italic line under it."""
        p = self.para(before=before)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT)
        self.text(p, left, bold=True)
        if right:
            self.text(p, "\t" + right, bold=True)
        if sub:
            p = self.para()
            self.text(p, sub, italic=True)

    def bullet(self, text: str):
        p = self.doc.add_paragraph(style="List Bullet")
        spacing(p, 0, 0)
        pf = p.paragraph_format
        pf.left_indent = Inches(0.22)
        pf.first_line_indent = Inches(-0.16)
        self.text(p, text)

    def labelled(self, label: str, value: str):
        p = self.para(after=1)
        self.text(p, f"{label}: ", bold=True)
        self.text(p, value)

    def plain(self, text: str, before=0):
        p = self.para(before=before)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self.text(p, text)

    def save(self, path: Path):
        self.doc.save(path)


# ------------------------------------------------------------------ one variant
def build_variant(content: dict, name: str, v: dict, out_dir: Path) -> Path:
    w = ResumeWriter()
    w.header(content["header"], v.get("headline", ""))

    w.section("Summary")
    w.plain(" ".join(v["summary"].split()))

    w.section("Skills")
    skills = content["skills"]
    for cat in v.get("skills") or list(skills):
        w.labelled(cat, ", ".join(skills[cat]))

    w.section("Experience")
    exp = content["experience"]
    for key in ("logicgo", "svnit"):
        e = exp[key]
        w.entry(f"{e['title']} – {e['org']}", e["dates"], before=4 if key == "svnit" else 2)
        for b in v.get(key) or list(e["bullets"]):
            w.bullet(e["bullets"][b])

    w.section("Projects")
    for i, key in enumerate(v.get("projects") or list(content["projects"])):
        pr = content["projects"][key]
        w.entry(pr["name"], pr.get("year", ""), sub=pr.get("tech", ""), before=2 if i == 0 else 4)
        for b in pr["bullets"]:
            w.bullet(b)

    w.section("Education")
    for ed in content["education"]:
        w.entry(f"{ed['degree']} – {ed['school']}", ed["dates"], before=2)

    certs = v.get("certifications") or []
    if certs:
        w.section("Certifications")
        for c in certs:
            w.bullet(content["certifications"][c])

    if content.get("languages"):
        p = w.para(before=5)
        w.text(p, "Languages: ", bold=True)
        w.text(p, content["languages"])

    path = out_dir / f"{name}.docx"
    w.save(path)
    return path


def pdf_page_count(pdf: Path) -> int:
    data = pdf.read_bytes()
    return len(re.findall(rb"/Type\s*/Page(?!s)", data))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Build role-specific resumes")
    ap.add_argument("variants", nargs="*", help="variant names (default: all)")
    ap.add_argument("--content", default=str(RES_DIR / "content.yaml"))
    ap.add_argument("--out", default=str(RES_DIR))
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args(argv)

    content = yaml.safe_load(Path(args.content).read_text(encoding="utf-8"))
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    names = args.variants or list(content["variants"])
    unknown = [n for n in names if n not in content["variants"]]
    if unknown:
        print("Unknown variants:", unknown, "| available:", list(content["variants"]))
        return 1

    built = [build_variant(content, n, content["variants"][n], out_dir) for n in names]
    for p in built:
        print("DOCX", p)
    if args.no_pdf:
        return 0
    try:
        from docx2pdf import convert
    except ImportError:
        print("Install docx2pdf (needs Microsoft Word) to create PDFs: pip install docx2pdf")
        return 1
    ok = True
    for p in built:
        pdf = p.with_suffix(".pdf")
        convert(str(p), str(pdf))
        pages = pdf_page_count(pdf)
        flag = "" if pages == 1 else f"   <-- {pages} pages, shorten content.yaml for a 1-page resume"
        ok &= pages == 1
        print(f"PDF  {pdf}  ({pages} page{'s' if pages != 1 else ''}){flag}")
    return 0 if ok else 2


if __name__ == "__main__":
    sys.exit(main())
