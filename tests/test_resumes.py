from pathlib import Path

import yaml
from docx import Document

import build_resumes
from naukri_bot.resumes import ResumePicker

ROOT = Path(__file__).resolve().parent.parent
EXAMPLE = ROOT / "resume_content.example.yaml"


def make_root(tmp_path, variants=("ai_ml_engineer", "genai_llm_engineer", "data_scientist")):
    res = tmp_path / "resumes"
    res.mkdir()
    content = yaml.safe_load(EXAMPLE.read_text(encoding="utf-8"))
    (res / "content.yaml").write_text(yaml.safe_dump(content), encoding="utf-8")
    for v in variants:
        (res / f"{v}.pdf").write_bytes(b"%PDF-1.4 test")
    return tmp_path


def test_picker_chooses_by_title(tmp_path):
    p = ResumePicker(make_root(tmp_path))
    assert p.choose("Generative AI Engineer") == "genai_llm_engineer"
    assert p.choose("Data Scientist - NLP") == "data_scientist"
    assert p.choose("AI Engineer") == "ai_ml_engineer"            # default = first variant
    assert p.choose("Barista") == "ai_ml_engineer"


def test_picker_copies_with_neutral_name(tmp_path):
    p = ResumePicker(make_root(tmp_path), upload_dir=tmp_path / "up")
    path = p.pick("LLM Engineer")
    assert path.exists() and path.parent.name == "genai_llm_engineer"
    assert path.name == "Your_Full_Name_Resume.pdf"


def test_picker_without_variants_uses_default(tmp_path):
    (tmp_path / "cv.pdf").write_bytes(b"%PDF-1.4")
    p = ResumePicker(tmp_path, "cv.pdf")
    assert p.choose("anything") is None and p.pick("anything") == tmp_path / "cv.pdf"


def test_build_variant_docx(tmp_path):
    content = yaml.safe_load(EXAMPLE.read_text(encoding="utf-8"))
    name = "genai_llm_engineer"
    path = build_resumes.build_variant(content, name, content["variants"][name], tmp_path)
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    for heading in ("SUMMARY", "SKILLS", "EXPERIENCE", "PROJECTS", "EDUCATION"):
        assert heading in text
    # the variant's own skill category comes first
    first_skill_line = next(p.text for p in doc.paragraphs if ":" in p.text and p.text.split(":")[0] in content["skills"])
    assert first_skill_line.startswith(content["variants"][name]["skills"][0])
    # header links are real hyperlinks
    assert any("hyperlink" in r.reltype for r in doc.part.rels.values())
